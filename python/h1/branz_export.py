"""
Get a chosen specification into the BRANZ H1/AS1 calculation workbook.

Our engine searches; BRANZ's workbook proves. A council accepts the BRANZ
spreadsheet because it is the tool the sector already agreed on, so the
deliverable is a filled copy of it rather than a report from us — which also
means nobody has to take our arithmetic on trust. `engine.py` stays a fast
pre-filter that makes the search possible; it is not what gets submitted.

Two ways across, because they carry very different licensing exposure:

`paste_blocks()` — the default. Returns tab-separated blocks the user pastes
into their own copy of the workbook, one per sheet. We never hold, process or
hand over BRANZ's file at all; the user downloads it from BRANZ themselves and
we hand them their own numbers. It also beats transcription on accuracy, which
is the thing that would otherwise put a transposed digit into a sealed consent.
The block boundaries are not arbitrary: the input columns are unlocked and the
first locked cell sits exactly where each block ends, so a paste at the anchor
can never collide with sheet protection.

`write()` — fills a template file directly. Better UX, but it means a copy of
BRANZ's workbook has to exist on the server, and handing filled copies to users
is redistribution of their work. Kept for when that has been cleared with BRANZ,
or for a flow where the user uploads their own copy.

Either way the sheets are not uniformly laid out, and assuming they were would
quietly put R-values in the wrong column of a compliance document:

    opaque elements   C description | D heated? | E area | F R-value
    joinery           C description | D ID | E area | F height | G width | H R-value

The workbook is BRANZ's, not ours, and is deliberately NOT committed here.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, List, Optional, Sequence

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from h1.tables import ELEMENT_TYPES, TableError

FIRST_ROW = 10
LAST_ROW = 59  # 50 input rows per sheet

# element type -> (sheet name, {field: column}). A missing "heated" key means the
# sheet has no heated-element column, which is the case for all the joinery.
LAYOUTS = {
    "slab_on_ground_floor": ("Slab Floors", {"description": "C", "heated": "D", "area": "E", "r_value": "F"}),
    "other_floor": ("Other Floors", {"description": "C", "heated": "D", "area": "E", "r_value": "F"}),
    "roof": ("Roof", {"description": "C", "heated": "D", "area": "E", "r_value": "F"}),
    "wall": ("Walls", {"description": "C", "heated": "D", "area": "E", "r_value": "F"}),
    "skylight": ("Skylights", {"description": "C", "area": "E", "r_value": "H"}),
    "glazing": ("Glazing (walls & doors)", {"description": "C", "area": "E", "r_value": "H"}),
    "opaque_door": ("Doors (opaque)", {"description": "C", "area": "E", "r_value": "H"}),
}

# Project Details: labels sit in column E, their inputs in column F.
PROJECT_CELLS = {
    "client": "F6",
    "project_name": "F7",
    "address": "F8",
    "designer": "F9",
    "date": "F10",
    "territorial_authority": "F12",
}

PROJECT_LABELS = {
    "client": "Client",
    "project_name": "Project name",
    "address": "Address",
    "designer": "Designer",
    "date": "Date",
    "territorial_authority": "Territorial Authority",
}

# Headers across each paste block, in column order from `description` to
# `r_value`. Columns with no field of ours behind them are pasted empty, which
# is what the workbook expects — they are the user's own joinery schedule
# details, not something we can know. `test_paste_headers_match_the_span`
# keeps these honest against LAYOUTS.
PASTE_HEADERS = {
    "slab_on_ground_floor": ["Description", "heated floor?", "Area (m²)", "R-value (m²·K/W)"],
    "other_floor": ["Description", "heated floor?", "Area (m²)", "R-value (m²·K/W)"],
    "roof": ["Description", "heated ceiling?", "Area (m²)", "R-value (m²·K/W)"],
    "wall": ["Description", "heated wall?", "Area (m²)", "R-value (m²·K/W)"],
    "skylight": ["Description", "ID", "Area (m²)", "Height (mm)", "Width (mm)", "R-value (m²·K/W)"],
    "glazing": ["Description", "ID", "Area (m²)", "Height (mm)", "Width (mm)", "R-value (m²·K/W)"],
    "opaque_door": ["Description", "ID", "Area (m²)", "Height (mm)", "Width (mm)", "R-value (m²·K/W)"],
}


@dataclass(frozen=True)
class Row:
    element_type: str
    description: str
    area_m2: float
    r_value: float
    heated: bool = False


@dataclass(frozen=True)
class ProjectDetails:
    project_name: str = ""
    client: str = ""
    address: str = ""
    designer: str = ""
    date: str = ""
    territorial_authority: str = ""


@dataclass(frozen=True)
class Block:
    """One sheet's worth of values, ready to paste at `anchor`."""

    sheet: str
    anchor: str          # e.g. "C10" — click this cell in Excel, then paste
    span: str            # e.g. "C10:F13" — what the paste will cover
    headers: List[str]
    rows: List[List[str]]
    tsv: str
    row_count: int


@dataclass(frozen=True)
class PasteInstructions:
    blocks: List[Block] = field(default_factory=list)
    project: List[dict] = field(default_factory=list)   # {cell, label, value}
    notes: List[str] = field(default_factory=list)


def paste_blocks(
    rows: Sequence[Row], project: Optional[ProjectDetails] = None
) -> PasteInstructions:
    """Turn a specification into blocks the user pastes into their own workbook.

    Excel splits tab-separated text across columns on paste, so one paste per
    sheet fills every value at once. That keeps BRANZ's file entirely in the
    user's hands while removing the transcription step, which is where a
    transposed digit would otherwise enter a sealed consent document.
    """
    _validate(rows)

    by_type = {}
    for row in rows:
        by_type.setdefault(row.element_type, []).append(row)

    blocks = []
    for element_type in ELEMENT_TYPES:
        sheet_rows = by_type.get(element_type)
        if not sheet_rows:
            continue
        blocks.append(_block(element_type, sheet_rows))

    details = []
    if project is not None:
        for field_name, cell in PROJECT_CELLS.items():
            value = getattr(project, field_name)
            if value:
                details.append({"cell": cell, "label": PROJECT_LABELS[field_name], "value": value})

    return PasteInstructions(
        blocks=blocks,
        project=details,
        notes=[
            "Download the H1/AS1 6th edition calculation tool from BRANZ, then for each "
            "block below click the anchor cell on that sheet and paste.",
            "Paste covers only the blue input columns — the workbook's own formulas and "
            "locked cells are never touched.",
        ],
    )


def _block(element_type: str, rows: Sequence[Row]) -> Block:
    sheet_name, columns = LAYOUTS[element_type]
    span_columns = _span(columns["description"], columns["r_value"])
    headers = PASTE_HEADERS[element_type]

    table = [_paste_row(row, columns, span_columns) for row in rows]
    last_row = FIRST_ROW + len(rows) - 1
    return Block(
        sheet=sheet_name,
        anchor=f"{span_columns[0]}{FIRST_ROW}",
        span=f"{span_columns[0]}{FIRST_ROW}:{span_columns[-1]}{last_row}",
        headers=headers,
        rows=table,
        tsv="\n".join("\t".join(cells) for cells in table),
        row_count=len(rows),
    )


def _paste_row(row: Row, columns: dict, span_columns: Sequence[str]) -> List[str]:
    by_column = {
        columns["description"]: _clean(row.description),
        columns["area"]: _number(row.area_m2),
        columns["r_value"]: _number(row.r_value),
    }
    if "heated" in columns:
        by_column[columns["heated"]] = "Yes" if row.heated else "No"
    return [by_column.get(letter, "") for letter in span_columns]


def _span(first: str, last: str) -> List[str]:
    start, end = column_index_from_string(first), column_index_from_string(last)
    return [get_column_letter(i) for i in range(start, end + 1)]


def _clean(text: str) -> str:
    """A tab or newline in a description would shift every later cell of the paste."""
    return " ".join(str(text).split())


def _number(value: float) -> str:
    rounded = round(float(value), 3)
    return str(int(rounded)) if rounded == int(rounded) else str(rounded)


def rows_from_candidate(candidate) -> List[Row]:
    """Turn an `optimise.Candidate` into workbook rows, keeping the assembly names."""
    return [
        Row(
            element_type=s.element_type,
            description=f"{s.assembly} — {s.name}",
            area_m2=s.area_m2,
            r_value=s.r_value,
        )
        for s in candidate.selections
    ]


def write(
    rows: Sequence[Row],
    template_path: Path,
    output_path: Path,
    project: Optional[ProjectDetails] = None,
) -> Path:
    _validate(rows)
    template_path, output_path = Path(template_path), Path(output_path)
    if not template_path.exists():
        raise TableError(
            f"No BRANZ template at {template_path}. Download the "
            f"'Calculation Method H1/AS1 6th Ed' workbook from BRANZ and point "
            f"H1_BRANZ_TEMPLATE at it."
        )

    workbook = openpyxl.load_workbook(template_path)
    missing = [name for name, _ in LAYOUTS.values() if name not in workbook.sheetnames]
    if missing:
        raise TableError(
            f"{template_path.name} is missing the sheet(s) {missing} — it does not look "
            f"like the H1/AS1 6th edition calculation workbook"
        )

    if project is not None:
        _write_project(workbook["Project Details"], project)

    by_type = {}
    for row in rows:
        by_type.setdefault(row.element_type, []).append(row)

    for element_type, sheet_rows in by_type.items():
        sheet_name, columns = LAYOUTS[element_type]
        _write_sheet(workbook[sheet_name], columns, sheet_rows)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def _validate(rows: Sequence[Row]) -> None:
    if not rows:
        raise TableError("No rows to export")

    counts = {}
    for row in rows:
        if row.element_type not in ELEMENT_TYPES:
            raise TableError(f"{row.description}: unknown element type {row.element_type!r}")
        if row.area_m2 <= 0:
            raise TableError(f"{row.description}: area must be positive, got {row.area_m2}")
        if row.r_value <= 0:
            raise TableError(f"{row.description}: R-value must be positive, got {row.r_value}")
        if row.heated and "heated" not in LAYOUTS[row.element_type][1]:
            raise TableError(
                f"{row.description}: the {LAYOUTS[row.element_type][0]} sheet has no heated-element "
                f"column, so `heated` cannot be recorded there"
            )
        counts[row.element_type] = counts.get(row.element_type, 0) + 1

    capacity = LAST_ROW - FIRST_ROW + 1
    over = {t: n for t, n in counts.items() if n > capacity}
    if over:
        raise TableError(
            f"The workbook has {capacity} rows per sheet; too many entries for "
            + ", ".join(f"{LAYOUTS[t][0]} ({n})" for t, n in over.items())
        )


def _write_project(sheet, project: ProjectDetails) -> None:
    for field, cell in PROJECT_CELLS.items():
        value = getattr(project, field)
        if value:
            sheet[cell] = value


def _write_sheet(sheet, columns: dict, rows: Iterable[Row]) -> None:
    """Clear the sheet's input range, then fill it. Formulas are never touched."""
    has_heated = "heated" in columns

    for row_number in range(FIRST_ROW, LAST_ROW + 1):
        sheet[f"{columns['description']}{row_number}"] = None
        sheet[f"{columns['area']}{row_number}"] = None
        sheet[f"{columns['r_value']}{row_number}"] = None
        if has_heated:
            # The template ships these as "No"; keep that rather than leaving a
            # blank the workbook would read as an error.
            sheet[f"{columns['heated']}{row_number}"] = "No"

    for offset, row in enumerate(rows):
        number = FIRST_ROW + offset
        sheet[f"{columns['description']}{number}"] = row.description
        sheet[f"{columns['area']}{number}"] = round(row.area_m2, 3)
        sheet[f"{columns['r_value']}{number}"] = round(row.r_value, 3)
        if has_heated:
            sheet[f"{columns['heated']}{number}"] = "Yes" if row.heated else "No"
